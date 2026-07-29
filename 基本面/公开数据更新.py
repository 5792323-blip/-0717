#!/usr/bin/env python3
"""用免费公开数据更新行业景气度台账。

默认只生成检查报告。传入 ``--apply`` 后，只有通过时效校验的数据才会写入工作簿；
评分仍由研究者按 v3 的评分口径复核，避免将单个宏观读数误当成行业结论。
"""

import argparse
import json
import shutil
from datetime import date, datetime
from pathlib import Path

import yaml
import pandas as pd
from openpyxl import load_workbook


根目录 = Path(__file__).resolve().parent
默认工作簿 = 根目录 / "全行业景气度量化跟踪体系_v3_全行业核心驱动版.xlsx"
默认映射 = 根目录 / "公开数据映射.yaml"


def 读取最新有效值(frame):
    """从 AKShare 宏观表中取最后一条有“今值”的记录。"""
    if frame is None or frame.empty or "日期" not in frame.columns or "今值" not in frame.columns:
        return None
    data = frame.copy()
    data["日期"] = data["日期"].astype(str).str[:10]
    data["今值"] = data["今值"].astype(str).str.replace(",", "", regex=False)
    data["数值"] = pd.to_numeric(data["今值"], errors="coerce")
    data = data.dropna(subset=["数值"]).sort_values("日期")
    if data.empty:
        return None
    row = data.iloc[-1]
    return {"日期": str(row["日期"]), "数值": float(row["数值"])}


def 数据是否新鲜(发布日期, 基准日, 最大时效天数):
    try:
        published = datetime.strptime(发布日期, "%Y-%m-%d").date()
        reference = datetime.strptime(基准日, "%Y-%m-%d").date()
    except ValueError:
        return False
    return published >= reference and (date.today() - published).days <= 最大时效天数


def 拉取公开宏观序列(映射项):
    import akshare as ak

    frame = getattr(ak, 映射项["函数"])()
    return 读取最新有效值(frame)


def 构建候选更新(映射, 台账):
    """返回候选更新及所有跳过原因，不在这里写工作簿。"""
    candidates, skipped = [], []
    for item in 映射.get("公开宏观序列", []):
        try:
            latest = 拉取公开宏观序列(item)
        except Exception as error:  # 网络或第三方公开接口失败时只记录。
            skipped.append({"匹配词": item["匹配词"], "原因": f"拉取失败: {type(error).__name__}"})
            continue
        if not latest:
            skipped.append({"匹配词": item["匹配词"], "原因": "无有效今值"})
            continue
        if not 数据是否新鲜(latest["日期"], 映射["工作簿数据基准日"], int(映射["最大时效天数"])):
            skipped.append({"匹配词": item["匹配词"], "原因": f"数据过期: {latest['日期']}"})
            continue
        for row in range(2, 台账.max_row + 1):
            indicator = str(台账.cell(row, 4).value or "")
            if item["匹配词"] in indicator:
                candidates.append({
                    "row": row,
                    "指标": indicator,
                    "数值": latest["数值"],
                    "日期": latest["日期"],
                    "单位": item.get("单位", ""),
                    "来源": f"AKShare/{item['函数']}",
                })
    return candidates, skipped


def 应用更新(台账, candidates):
    for item in candidates:
        row = item["row"]
        台账.cell(row, 6, item["数值"])
        台账.cell(row, 7, item["单位"] or 台账.cell(row, 7).value)
        台账.cell(row, 10, item["日期"])
        台账.cell(row, 11, "待更新")
        台账.cell(row, 14, f"自动更新：{item['来源']}；请按评分口径复核评分。")


def main():
    parser = argparse.ArgumentParser(description="公开数据行业景气度更新器")
    parser.add_argument("--workbook", type=Path, default=默认工作簿)
    parser.add_argument("--mapping", type=Path, default=默认映射)
    parser.add_argument("--apply", action="store_true", help="通过时效校验后写入工作簿")
    args = parser.parse_args()

    with args.mapping.open(encoding="utf-8") as source:
        mapping = yaml.safe_load(source)
    workbook = load_workbook(args.workbook)
    ledger = workbook["指标更新台账_v3"]
    candidates, skipped = 构建候选更新(mapping, ledger)
    report = {
        "运行时间": datetime.now().isoformat(timespec="seconds"),
        "工作簿": str(args.workbook),
        "模式": "apply" if args.apply else "dry-run",
        "候选更新": candidates,
        "跳过": skipped,
    }
    if args.apply and candidates:
        backup_dir = 根目录 / "备份"
        backup_dir.mkdir(exist_ok=True)
        backup = backup_dir / f"{args.workbook.stem}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        shutil.copy2(args.workbook, backup)
        应用更新(ledger, candidates)
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(args.workbook)
        report["备份"] = str(backup)
    report_dir = 根目录 / "运行记录"
    report_dir.mkdir(exist_ok=True)
    report_path = report_dir / "最新公开数据更新报告.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
