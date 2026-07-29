from openpyxl import Workbook

from 基本面.公开数据更新 import 应用更新, 数据是否新鲜, 读取最新有效值


def test_reads_last_non_null_macro_value():
    import pandas as pd

    frame = pd.DataFrame({
        "日期": ["2026-05-01", "2026-06-01", "2026-07-01"],
        "今值": ["8.1", None, "8.4"],
    })
    assert 读取最新有效值(frame) == {"日期": "2026-07-01", "数值": 8.4}


def test_rejects_stale_or_pre_baseline_data():
    assert 数据是否新鲜("2026-06-01", "2026-05-31", 365)
    assert not 数据是否新鲜("2026-05-30", "2026-05-31", 365)


def test_apply_updates_value_but_preserves_manual_score():
    ledger = Workbook().active
    ledger.cell(2, 9, 7)
    应用更新(ledger, [{
        "row": 2,
        "数值": 52.1,
        "日期": "2026-07-28",
        "单位": "点",
        "来源": "AKShare/test",
    }])

    assert ledger.cell(2, 6).value == 52.1
    assert ledger.cell(2, 9).value == 7
    assert ledger.cell(2, 10).value == "2026-07-28"
